#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Лаб. 3.2.5, п. 2.2 и 2.7.1: период свободных колебаний в зависимости от ёмкости.

Данные берутся прямо из таблицы 325.xlsx (лист «1. Свободные колебания», блок «2. Зависимость периода
свободных колебаний от емкости»: C магазина, мкФ  и  T_эксп, мкс).

Теория:  T = 2π·√(L·(C + C0)),  C0 - нулевая ёмкость контура (монтаж + магазин + ёмкость соединительной
коробки), определяется при C = 0 из периода T0 (п. 2.2.4): C0 = T0²/(4π²L).
Кроме того C0 (и L) определяются по всем точкам: T² = 4π²L·C + 4π²L·C0  - прямая в координатах (C, T²),
МНК с учётом погрешностей по обеим осям. Полученная L сравнивается с измеренной LCR-метром.

Запуск:  python analyze_period.py [--xlsx 325.xlsx] [--sigT 1.0] [--Crel 0.01] [--no-show]
"""
import argparse
import os
import sys

import numpy as np
import matplotlib
import matplotlib.pyplot as plt

plt.rcParams.update({"font.size": 12, "axes.titlesize": 14, "axes.labelsize": 13, "legend.fontsize": 11})

L_LCR, L_LCR_SIG = 99.944e-3, 0.1e-3     # Гн: измерено LCR-метром GW Instek LCR-7819 (50/500/1500 Гц)
SHEET = "1. Свободные колебания"

DATA_DIR = r"C:\Utilities\Projects\MIPT_labs\lab_3.2.4+5\data"
FIGURES_DIR = r"C:\Utilities\Projects\MIPT_labs\lab_3.2.4+5\figures"


def read_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET]
    C, T = [], []
    start = None
    for r in range(1, ws.max_row + 1):                       # заголовок блока ищем по тексту
        v = ws.cell(r, 3).value
        if isinstance(v, str) and v.startswith("T_эксп"):
            start = r + 1
            break
    if start is None:
        sys.exit("Не найден блок с T_эксп на листе «1. Свободные колебания»")
    r = start
    while isinstance(ws.cell(r, 2).value, (int, float)) and isinstance(ws.cell(r, 3).value, (int, float)):
        C.append(float(ws.cell(r, 2).value) * 1e-6)           # мкФ -> Ф
        T.append(float(ws.cell(r, 3).value) * 1e-6)           # мкс -> с
        r += 1
    return np.array(C), np.array(T)


def wls_line(x, y, sx, sy, iters=10):
    """y = a + b x, эффективная дисперсия (погрешности по x и y). Возвращает a, b, cov(a,b), chi2/ndf."""
    b = np.polyfit(x, y, 1)[0]
    for _ in range(iters):
        w = 1.0 / (sy ** 2 + (b * sx) ** 2)
        A = np.vstack([np.ones_like(x), x]).T
        cov = np.linalg.inv(A.T @ (A * w[:, None]))
        a, b = cov @ (A.T @ (w * y))
    chi2 = np.sum(w * (y - a - b * x) ** 2) / max(len(x) - 2, 1)
    return a, b, cov, chi2


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=os.path.join(DATA_DIR, "325.xlsx"))
    ap.add_argument("--sigT", type=float, default=1.0, help="погрешность отсчёта T курсорами, мкс (ОЦЕНКА)")
    ap.add_argument("--Crel", type=float, default=0.01, help="отн. погрешность ёмкости магазина (ОЦЕНКА)")
    ap.add_argument("--out", default=os.path.join(FIGURES_DIR, "period"))
    ap.add_argument("--no-show", action="store_true")
    args = ap.parse_args()
    if args.no_show:
        matplotlib.use("Agg")

    C, T = read_xlsx(args.xlsx)
    sT = np.full_like(T, args.sigT * 1e-6)
    sC = args.Crel * C + 0.0
    n = len(C)

    # ---- 1) C0 по нулевой точке (п. 2.2.4) ----
    k = 4 * np.pi ** 2
    C0_zero = T[0] ** 2 / (k * L_LCR)
    sC0_zero = C0_zero * np.hypot(2 * sT[0] / T[0], L_LCR_SIG / L_LCR)

    # ---- 2) прямая T² = a + b·C  по всем точкам (L и C0 свободны) ----
    y, sy = T ** 2, 2 * T * sT
    a, b, cov, chi2 = wls_line(C, y, sC, sy)
    L_fit = b / k
    C0_fit = a / b
    sb, sa = np.sqrt(cov[1, 1]), np.sqrt(cov[0, 0])
    sL_fit = sb / k * np.sqrt(max(chi2, 1.0))
    sC0_fit = C0_fit * np.sqrt((sa / a) ** 2 + (sb / b) ** 2 - 2 * cov[0, 1] / (a * b)) * np.sqrt(max(chi2, 1.0))

    # ---- 3) C0 при L, фиксированном по LCR-метру (взвешенное среднее по точкам) ----
    C0_pts = T ** 2 / (k * L_LCR) - C
    sC0_pts = np.sqrt((2 * T * sT / (k * L_LCR)) ** 2 + sC ** 2 + (C0_pts + C) ** 2 * (L_LCR_SIG / L_LCR) ** 2)
    w = 1 / sC0_pts ** 2
    C0_w = np.sum(w * C0_pts) / np.sum(w)
    sC0_w = 1 / np.sqrt(np.sum(w))
    chi2_w = np.sum(w * (C0_pts - C0_w) ** 2) / (n - 1)

    # ---- теория для графика: L по LCR, C0 из шага 3 ----
    T_th = 2 * np.pi * np.sqrt(L_LCR * (C + C0_w))
    # погрешность T_теор: от C (магазин), C0, L
    sT_th = T_th * 0.5 * np.sqrt((np.hypot(sC, sC0_w) / (C + C0_w)) ** 2 + (L_LCR_SIG / L_LCR) ** 2)
    resid = T - T_th
    sres = np.hypot(sT, sT_th)

    print(f"Точек: {n};  σ(T) = {args.sigT:.2f} мкс, σ(C) = {100*args.Crel:.1f} % (оценки)")
    print(f"L (LCR-метр)        = {L_LCR*1e3:.3f} ± {L_LCR_SIG*1e3:.2f} мГн")
    print(f"C0 по T(C=0), п.2.2.4 = {C0_zero*1e9:.3f} ± {sC0_zero*1e9:.3f} нФ")
    print(f"C0 по всем точкам, L из LCR = {C0_w*1e9:.3f} ± {sC0_w*1e9:.3f} нФ (χ²/ndf = {chi2_w:.2f})")
    print(f"Прямая T² = a + b·C:  L = b/4π² = {L_fit*1e3:.2f} ± {sL_fit*1e3:.2f} мГн, "
          f"C0 = a/b = {C0_fit*1e9:.3f} ± {sC0_fit*1e9:.3f} нФ  (χ²/ndf = {chi2:.2f})")
    print(f"  L (фит) / L (LCR) - 1 = {100*(L_fit/L_LCR-1):+.2f} %")
    print("\n  C, нФ   T_эксп, мкс   T_теор, мкс   T_эксп - T_теор, мкс   в σ")
    for i in range(n):
        print(f"  {C[i]*1e9:5.1f}   {T[i]*1e6:9.1f}   {T_th[i]*1e6:10.2f}   {resid[i]*1e6:12.2f}   {resid[i]/sres[i]:+8.2f}")
    print(f"\nRMS отклонения: {np.sqrt(np.mean(resid**2))*1e6:.2f} мкс "
          f"({100*np.sqrt(np.mean((resid/T_th)**2)):.2f} % от T)")

    # ---- графики ----
    fig, ax = plt.subplots(1, 3, figsize=(17, 5.6))
    a1 = ax[0]
    lim = [T.min() * 0.9e6, T.max() * 1.05e6]
    a1.plot(lim, lim, "k--", lw=1, label="$T_{эксп} = T_{теор}$")
    a1.errorbar(T_th * 1e6, T * 1e6, xerr=2 * sT_th * 1e6, yerr=2 * sT * 1e6, fmt="o", color="C0", mec="k",
                ecolor="k", elinewidth=1.2, capsize=3, label="эксперимент (±2σ)")
    a1.set_xlim(lim); a1.set_ylim(lim); a1.set_aspect("equal")
    a1.set_xlabel("$T_{теор} = 2\\pi\\sqrt{L(C+C_0)}$, мкс"); a1.set_ylabel("$T_{эксп}$, мкс")
    a1.set_title("$T_{эксп} = f(T_{теор})$ (п. 2.7.1)")
    a1.text(0.03, 0.97, f"L = {L_LCR*1e3:.2f} мГн (LCR)\n$C_0$ = {C0_w*1e9:.2f} ± {sC0_w*1e9:.2f} нФ",
            transform=a1.transAxes, va="top", fontsize=11, bbox=dict(fc="w", ec="0.7"))
    a1.legend(loc="lower right"); a1.grid(True, ls="--", alpha=0.5)

    a2 = ax[1]
    a2.errorbar(C * 1e9, y * 1e12, xerr=2 * sC * 1e9, yerr=2 * sy * 1e12, fmt="o", color="C0", mec="k",
                ecolor="k", elinewidth=1.2, capsize=3, label="эксперимент (±2σ)")
    cc = np.linspace(-C0_fit * 1.2, C.max() * 1.05, 50)
    a2.plot(cc * 1e9, (a + b * cc) * 1e12, "C3-", label=f"МНК: L = {L_fit*1e3:.1f} ± {sL_fit*1e3:.1f} мГн")
    a2.axvline(0, color="gray", lw=0.7)
    a2.plot([-C0_fit * 1e9], [0], "rx", ms=9, mew=2, label=f"$-C_0$ = {-C0_fit*1e9:.2f} ± {sC0_fit*1e9:.2f} нФ")
    a2.set_xlabel("C магазина, нФ"); a2.set_ylabel("$T^2$, мкс$^2$")
    a2.set_title("Линеаризация: $T^2 = 4\\pi^2 L (C + C_0)$")
    a2.legend(loc="upper left", fontsize=10); a2.grid(True, ls="--", alpha=0.5)

    a3 = ax[2]
    a3.axhline(0, color="k", lw=0.8)
    a3.errorbar(C * 1e9, resid * 1e6, yerr=2 * sres * 1e6, fmt="o", color="C0", mec="k", ecolor="k",
                elinewidth=1.2, capsize=3)
    a3.set_xlabel("C магазина, нФ"); a3.set_ylabel("$T_{эксп} - T_{теор}$, мкс")
    a3.set_title("Отклонение от теории (±2σ)"); a3.grid(True, ls="--", alpha=0.5)

    fig.tight_layout()
    fig.savefig(args.out + ".pdf")
    print(f"\nСохранено: {args.out}.pdf")
    if not args.no_show:
        plt.show()


if __name__ == "__main__":
    main()
